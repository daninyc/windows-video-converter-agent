from pathlib import Path
from zipfile import BadZipFile

import pytest
from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from media_converter.common.errors import AppError
from media_converter.config import AppPaths
from media_converter.devtools.fixtures import generate_smoke_fixtures, safe_fixture_path
from tests.fixture_support import load_fixture_manifest

pytestmark = pytest.mark.integration


def test_real_media_documents_manifest_and_idempotence(tmp_path):
    root = tmp_path / "fixtures"
    paths = AppPaths.from_argument(tmp_path / "data")
    manifest = generate_smoke_fixtures(root, paths)
    assert manifest == load_fixture_manifest(root)
    assert generate_smoke_fixtures(root, paths) == manifest
    media = [item for item in manifest["files"] if item["kind"] == "media"]
    assert {".mp4", ".mov", ".mkv", ".avi", ".webm"} <= {
        Path(item["path"]).suffix for item in media
    }
    for item in media:
        if item["expected"].get("probe_error"):
            continue
        lo, hi = item["expected"]["duration_us_range"]
        assert lo <= item["observed"]["duration_us"] <= hi
        assert [stream["time_base"] for stream in item["observed"]["streams"]] == item["expected"][
            "time_bases"
        ]
        assert "{fixture_root}" in " ".join(item["command"])
    pdf = PdfReader(root / "documents/说明.pdf")
    assert "H.264" in pdf.pages[0].extract_text()
    assert not PdfReader(root / "documents/无文字.pdf").pages[0].extract_text()
    encrypted = PdfReader(root / "documents/加密.pdf")
    assert encrypted.is_encrypted
    assert encrypted.decrypt("fixture-only")
    assert "AAC" in encrypted.pages[0].extract_text()
    assert "自生成测试文档" == Document(root / "documents/说明.docx").paragraphs[0].text
    with pytest.raises(PdfReadError):
        PdfReader(root / "documents/损坏.pdf")
    with pytest.raises(BadZipFile):
        Document(root / "documents/损坏.docx")
    (root / "documents/说明.txt").write_text("changed", "utf-8")
    with pytest.raises(AppError):
        generate_smoke_fixtures(root, paths)


def test_refuses_foreign_files_and_path_escape(tmp_path):
    foreign = tmp_path / "mine.txt"
    foreign.write_text("keep me", "utf-8")
    with pytest.raises(AppError):
        generate_smoke_fixtures(tmp_path, AppPaths.from_argument(tmp_path / "data"))
    assert foreign.read_text("utf-8") == "keep me"
    for relative in ("../escape.mp4", "D:/escape.mp4", "C:\\escape.pdf", "sub/../../escape"):
        with pytest.raises(AppError):
            safe_fixture_path(tmp_path, relative)
