import hashlib
import json
import os
import tempfile
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path, PureWindowsPath

from media_converter.common.errors import AppError, ErrorCode
from media_converter.config import AppPaths
from media_converter.media.toolchain import inspect_toolchain, run_tool

MANIFEST_NAME = "manifest.json"
GENERATOR_ID = "media-converter-smoke-v1"
EXPECTED_TIME_BASES = {
    "media/中文 空格/基础.mp4": ["1/12288", "1/48000"],
    "media/基础.mov": ["1/12288", "1/48000"],
    "media/基础.avi": ["1/12", "1/48000"],
    "media/基础.webm": ["1/1000", "1/1000"],
    "media/无音频.mp4": ["1/12288"],
    "media/双音轨 双字幕.mkv": ["1/1000"] * 5,
    "media/兼容字幕.mp4": ["1/12288", "1/48000", "1/1000000"],
    "media/旋转.mov": ["1/12288", "1/48000"],
    "media/变帧率.mp4": ["1/12288"],
    "media/超短.mp4": ["1/12800"],
}


def safe_fixture_path(root: Path, relative: str) -> Path:
    windows = PureWindowsPath(relative)
    if windows.drive or windows.is_absolute() or ".." in windows.parts:
        raise AppError(ErrorCode.ACCESS_DENIED)
    path = (root / relative).resolve()
    if not path.is_relative_to(root.resolve()) or path == root.resolve():
        raise AppError(ErrorCode.ACCESS_DENIED)
    return path


def sha256_file(path: Path) -> str:
    with path.open("rb") as file:
        return hashlib.file_digest(file, "sha256").hexdigest()


def verify_manifest(root: Path) -> dict:
    try:
        manifest = json.loads((root / MANIFEST_NAME).read_text("utf-8"))
        if manifest.get("generator") != GENERATOR_ID:
            raise ValueError("foreign manifest")
        for item in manifest["files"]:
            path = safe_fixture_path(root, item["path"])
            if not path.is_file() or sha256_file(path) != item["sha256"]:
                raise ValueError("fixture changed")
        return manifest
    except (OSError, ValueError, KeyError, TypeError) as error:
        raise AppError(ErrorCode.FILE_CHANGED) from error


def _probe(binary: Path, path: Path) -> dict:
    return json.loads(
        run_tool(binary, ["-v", "error", "-show_streams", "-show_format", "-of", "json", str(path)])
    )


def _make_documents(stage: Path) -> list[dict]:
    from docx import Document
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    documents = []
    text = "自生成测试文字。视频编码 H.264，音频 AAC。\nSecond paragraph: fixture evidence.\n"
    for name in ("说明.txt", "说明.md", "运行.log"):
        path = stage / "documents" / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, "utf-8")
        documents.append(
            {
                "path": path.relative_to(stage).as_posix(),
                "kind": "document",
                "expected": {"encoding": "utf-8", "text": text},
            }
        )
    document = Document()
    document.add_paragraph("自生成测试文档")
    document.add_paragraph("Video codec H.264; audio codec AAC.")
    document.save(stage / "documents" / "说明.docx")
    documents.append(
        {
            "path": "documents/说明.docx",
            "kind": "document",
            "expected": {"paragraphs": ["自生成测试文档", "Video codec H.264; audio codec AAC."]},
        }
    )
    for name, has_text, encrypted in [
        ("说明.pdf", True, False),
        ("无文字.pdf", False, False),
        ("加密.pdf", True, True),
    ]:
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        if has_text:
            font = DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                }
            )
            page[NameObject("/Resources")] = DictionaryObject(
                {
                    NameObject("/Font"): DictionaryObject(
                        {NameObject("/F1"): writer._add_object(font)}
                    ),
                }
            )
            stream = DecodedStreamObject()
            stream.set_data(b"BT /F1 12 Tf 72 700 Td (Fixture video H.264 audio AAC.) Tj ET")
            page[NameObject("/Contents")] = writer._add_object(stream)
        if encrypted:
            writer.encrypt("fixture-only", algorithm="AES-256")
        writer.write(stage / "documents" / name)
        documents.append(
            {
                "path": "documents/" + name,
                "kind": "document",
                "expected": {
                    "pages": 1,
                    "encrypted": encrypted,
                    "text": "Fixture video H.264 audio AAC." if has_text else "",
                },
            }
        )
    for name, content in [
        ("损坏.pdf", b"%PDF-1.7\ncorrupt fixture"),
        ("损坏.docx", b"not a zip archive"),
    ]:
        (stage / "documents" / name).write_bytes(content)
        documents.append(
            {"path": "documents/" + name, "kind": "document", "expected": {"parse_error": True}}
        )
    return documents


def generate_smoke_fixtures(root: Path, paths: AppPaths) -> dict:
    root = root.resolve()
    root.mkdir(parents=True, exist_ok=True)
    if (root / MANIFEST_NAME).exists():
        return verify_manifest(root)
    # Refuse a mixed directory, including abandoned runs; never overwrite the user's files.
    if any(root.iterdir()):
        raise AppError(ErrorCode.ACCESS_DENIED)
    snapshot = inspect_toolchain(paths)
    snapshot.require_ready()
    ffmpeg, ffprobe = snapshot.ffmpeg.path, snapshot.ffprobe.path
    records = []
    with tempfile.TemporaryDirectory(prefix=".fixture-build-", dir=root) as temporary:
        stage = Path(temporary).resolve()
        if not stage.is_relative_to(root):
            raise AppError(ErrorCode.ACCESS_DENIED)

        def path(relative):
            target = safe_fixture_path(stage, relative)
            target.parent.mkdir(parents=True, exist_ok=True)
            return target

        def media(relative, arguments, codecs, **expected):
            destination = path(relative)
            command = [
                "-hide_banner",
                "-loglevel",
                "error",
                "-nostdin",
                "-n",
                *arguments,
                str(destination),
            ]
            run_tool(ffmpeg, command, timeout=60)
            probe = _probe(ffprobe, destination)
            streams = probe["streams"]
            actual_codecs = [stream["codec_name"] for stream in streams]
            time_bases = [stream["time_base"] for stream in streams]
            if actual_codecs != codecs or time_bases != EXPECTED_TIME_BASES[relative]:
                raise AppError(ErrorCode.PROCESS_FAILED)
            duration_us = int(Decimal(probe["format"]["duration"]) * 1_000_000)
            if duration_us <= 0:
                raise AppError(ErrorCode.PROCESS_FAILED)
            item = {
                "path": relative,
                "kind": "media",
                "command": [
                    str(ffmpeg),
                    *[arg.replace(str(stage), "{fixture_root}") for arg in command],
                ],
                "expected": {
                    "codecs": codecs,
                    "time_bases": EXPECTED_TIME_BASES[relative],
                    "width": 160,
                    "height": 90,
                    "duration_us_range": [900_000, 1_200_000],
                    **expected,
                },
                "observed": {
                    "duration_us": duration_us,
                    "streams": streams,
                    "format": probe["format"]["format_name"],
                },
            }
            video = next(stream for stream in streams if stream["codec_type"] == "video")
            if (video["width"], video["height"]) != (160, 90):
                raise AppError(ErrorCode.PROCESS_FAILED)
            low, high = item["expected"]["duration_us_range"]
            if not low <= duration_us <= high:
                raise AppError(ErrorCode.PROCESS_FAILED)
            if "audio_languages" in expected:
                audio = [stream for stream in streams if stream["codec_type"] == "audio"]
                if [stream.get("tags", {}).get("language") for stream in audio] != expected[
                    "audio_languages"
                ]:
                    raise AppError(ErrorCode.PROCESS_FAILED)
                if [stream["disposition"]["default"] for stream in audio] != expected[
                    "audio_default_flags"
                ]:
                    raise AppError(ErrorCode.PROCESS_FAILED)
            if "frame_rate" in expected:
                rate = expected["frame_rate"]
                if video["r_frame_rate"] != f"{rate['numerator']}/{rate['denominator']}":
                    raise AppError(ErrorCode.PROCESS_FAILED)
            if "rotation_deg" in expected:
                rotations = [entry.get("rotation") for entry in video.get("side_data_list", [])]
                if expected["rotation_deg"] not in rotations:
                    raise AppError(ErrorCode.PROCESS_FAILED)
            if expected.get("variable_frame_rate"):
                frames = json.loads(
                    run_tool(
                        ffprobe,
                        [
                            "-v",
                            "error",
                            "-select_streams",
                            "v:0",
                            "-show_frames",
                            "-show_entries",
                            "frame=best_effort_timestamp_time",
                            "-of",
                            "json",
                            str(destination),
                        ],
                    )
                )["frames"]
                times = [Decimal(frame["best_effort_timestamp_time"]) for frame in frames]
                deltas = {round(b - a, 4) for a, b in zip(times, times[1:], strict=False)}
                if len(deltas) < 2:
                    raise AppError(ErrorCode.PROCESS_FAILED)
                item["observed"]["frame_timestamps_seconds"] = [str(value) for value in times]
            records.append(item)
            return destination

        base = media(
            "media/中文 空格/基础.mp4",
            [
                "-f",
                "lavfi",
                "-i",
                "testsrc2=size=160x90:rate=12:duration=1",
                "-f",
                "lavfi",
                "-i",
                "sine=frequency=880:sample_rate=48000:duration=1",
                "-c:v",
                "libx264",
                "-preset",
                "ultrafast",
                "-threads",
                "1",
                "-pix_fmt",
                "yuv420p",
                "-c:a",
                "aac",
                "-shortest",
            ],
            ["h264", "aac"],
            frame_rate={"numerator": 12, "denominator": 1},
        )
        media("media/基础.mov", ["-i", str(base), "-map", "0", "-c", "copy"], ["h264", "aac"])
        media(
            "media/基础.avi",
            ["-i", str(base), "-c:v", "mpeg4", "-q:v", "4", "-threads", "1", "-c:a", "pcm_s16le"],
            ["mpeg4", "pcm_s16le"],
        )
        media(
            "media/基础.webm",
            [
                "-i",
                str(base),
                "-c:v",
                "libvpx-vp9",
                "-crf",
                "32",
                "-b:v",
                "0",
                "-threads",
                "1",
                "-c:a",
                "libopus",
            ],
            ["vp9", "opus"],
        )
        media("media/无音频.mp4", ["-i", str(base), "-an", "-c:v", "copy"], ["h264"])
        srt = path("subtitles/文字.srt")
        srt.write_text("1\n00:00:00,000 --> 00:00:00,900\nSynthetic subtitle\n", "utf-8")
        ass = path("subtitles/样式.ass")
        ass.write_text(
            "[Script Info]\nScriptType: v4.00+\nPlayResX: 160\nPlayResY: 90\n"
            "[V4+ Styles]\nFormat: Name, Fontname, Fontsize, PrimaryColour, SecondaryColour, "
            "OutlineColour, BackColour, Bold, Italic, Underline, StrikeOut, ScaleX, ScaleY, "
            "Spacing, Angle, BorderStyle, Outline, Shadow, Alignment, "
            "MarginL, MarginR, MarginV, Encoding\n"
            "Style: Default,Arial,14,&H00FFFFFF,&H000000FF,&H00000000,&H00000000,0,0,0,0,100,100,"
            "0,0,1,1,0,2,10,10,10,1\n[Events]\nFormat: Layer, Start, End, Style, Name, MarginL, "
            "MarginR, MarginV, Effect, Text\n"
            "Dialogue: 0,0:00:00.00,0:00:00.90,Default,,0,0,0,,Styled fixture\n",
            "utf-8",
        )
        media(
            "media/双音轨 双字幕.mkv",
            [
                "-i",
                str(base),
                "-f",
                "lavfi",
                "-i",
                "sine=frequency=440:sample_rate=48000:duration=1",
                "-i",
                str(srt),
                "-i",
                str(ass),
                "-map",
                "0:v:0",
                "-map",
                "0:a:0",
                "-map",
                "1:a:0",
                "-map",
                "2:0",
                "-map",
                "3:0",
                "-c:v",
                "copy",
                "-c:a",
                "aac",
                "-c:s",
                "copy",
                "-metadata:s:a:0",
                "language=zho",
                "-metadata:s:a:1",
                "language=eng",
                "-disposition:a:0",
                "default",
                "-disposition:a:1",
                "0",
            ],
            ["h264", "aac", "aac", "subrip", "ass"],
            audio_languages=["zho", "eng"],
            audio_default_flags=[1, 0],
            subtitle_compatibility={"matroska": ["subrip", "ass"], "mp4_incompatible": ["ass"]},
        )
        media(
            "media/兼容字幕.mp4",
            [
                "-i",
                str(base),
                "-i",
                str(srt),
                "-map",
                "0",
                "-map",
                "1:0",
                "-c",
                "copy",
                "-c:s",
                "mov_text",
            ],
            ["h264", "aac", "mov_text"],
        )
        media(
            "media/旋转.mov",
            ["-display_rotation:v:0", "90", "-i", str(base), "-map", "0", "-c", "copy"],
            ["h264", "aac"],
            rotation_deg=90,
        )
        media(
            "media/变帧率.mp4",
            [
                "-i",
                str(base),
                "-an",
                "-vf",
                "select=not(eq(mod(n\\,3)\\,1))",
                "-fps_mode",
                "vfr",
                "-c:v",
                "libx264",
                "-threads",
                "1",
            ],
            ["h264"],
            variable_frame_rate=True,
        )
        media(
            "media/超短.mp4",
            [
                "-f",
                "lavfi",
                "-i",
                "color=c=blue:size=160x90:rate=25",
                "-frames:v",
                "1",
                "-c:v",
                "libx264",
                "-threads",
                "1",
            ],
            ["h264"],
            duration_us_range=[40_000, 40_000],
        )
        broken = path("media/损坏.mp4")
        broken.write_bytes(b"deliberately invalid synthetic media")
        try:
            _probe(ffprobe, broken)
        except AppError as error:
            if error.info.code != ErrorCode.PROCESS_FAILED:
                raise
        else:
            raise AppError(ErrorCode.PROCESS_FAILED)
        records.append(
            {"path": "media/损坏.mp4", "kind": "media", "expected": {"probe_error": True}}
        )
        records += _make_documents(stage)
        records += [
            {"path": "subtitles/文字.srt", "kind": "subtitle_source"},
            {"path": "subtitles/样式.ass", "kind": "subtitle_source"},
        ]
        for item in records:
            file = safe_fixture_path(stage, item["path"])
            item["sha256"] = sha256_file(file)
            item["size_bytes"] = file.stat().st_size
        manifest = {
            "schema_version": 1,
            "generator": GENERATOR_ID,
            "created_at": datetime.now(UTC).isoformat(),
            "profile": "smoke",
            "license": "CC0-1.0",
            "provenance": "Self-generated lavfi pictures, tones and fixture text",
            "toolchain": snapshot.model_dump(mode="json"),
            "files": records,
        }
        for item in records:
            target = safe_fixture_path(root, item["path"])
            target.parent.mkdir(parents=True, exist_ok=True)
            # Atomic no-replace publication on the same volume; existing names always fail.
            os.link(safe_fixture_path(stage, item["path"]), target)
        with (root / MANIFEST_NAME).open("x", encoding="utf-8") as file:
            json.dump(manifest, file, ensure_ascii=False, indent=2)
    return manifest
